"""
Expense tracker with a monthly budget.

  "Spent 120 on lunch" / "Paid 45 for an auto" / "250 rupees recharge"
  "How much did I spend this week?" / "How much on food this month?"
  "Set my monthly budget to 6000"
  "Undo the last expense"
  "Export my expenses" (Excel file in exports/, opened for you)

Jarvis warns you when the month crosses 80% and 100% of the budget, and the
evening and weekly calls include what you spent.
"""
from __future__ import annotations

import os
import platform
import re
import subprocess
from collections import defaultdict
from datetime import date, timedelta

from plugins import _bus as bus
from plugins import _college_store as college
from plugins import _store as st

NAME = "expenses"
DEFAULT = {"items": [], "budget": 0, "warned": {}}
EXPORT_DIR = st.BASE_DIR / "exports"

CATEGORIES = {
    "food": ["lunch", "dinner", "breakfast", "snack", "tea", "coffee", "canteen", "mess", "swiggy",
             "zomato", "food", "juice", "biryani", "pizza", "burger", "dosa", "meal", "chai", "groceries", "grocery"],
    "travel": ["auto", "bus", "uber", "ola", "rapido", "metro", "train", "cab", "petrol", "fuel", "taxi", "ticket", "travel"],
    "education": ["book", "xerox", "print", "fees", "fee", "stationery", "pen", "course", "notes", "lab"],
    "bills": ["recharge", "wifi", "internet", "electricity", "rent", "bill", "phone", "subscription"],
    "fun": ["movie", "netflix", "spotify", "game", "party", "outing", "concert", "prime"],
    "shopping": ["shirt", "clothes", "shoes", "amazon", "flipkart", "myntra", "shopping", "gift"],
    "health": ["medicine", "pharmacy", "doctor", "gym"],
}

PLUGIN = {
    "name": "expenses",
    "description": (
        "Personal expense tracker in rupees. Use when the user says they spent/paid money "
        "('spent 120 on lunch'), asks how much they spent (today/week/month, optionally per "
        "category), sets a monthly budget, undoes the last expense, or wants an Excel export."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "action": {"type": "STRING", "description": "add, summary, set_budget, undo_last, export"},
            "amount": {"type": "NUMBER", "description": "Amount in rupees"},
            "category": {"type": "STRING",
                         "description": "food, travel, education, bills, fun, shopping, health or other"},
            "note": {"type": "STRING", "description": "What it was for, e.g. 'lunch at canteen'"},
            "date": {"type": "STRING", "description": "When (default today)"},
            "period": {"type": "STRING", "description": "today, yesterday, week, month, last_month"},
        },
        "required": ["action"],
    },
}


def run(parameters: dict, player=None, session_memory=None) -> str:
    try:
        a = (parameters.get("action") or "summary").lower().strip()
        if a == "add":
            return _add(parameters)
        if a == "set_budget":
            return _set_budget(parameters)
        if a == "undo_last":
            return _undo()
        if a == "export":
            return _export(parameters)
        return _summary(parameters.get("period") or "month", parameters.get("category") or "")
    except Exception as e:
        return f"Sir, the expense tracker hit an error: {e}"


def _rs(x: float) -> str:
    return f"₹{x:,.0f}" if float(x).is_integer() else f"₹{x:,.2f}"


def _guess_category(note: str) -> str:
    n = (note or "").lower()
    for cat, words in CATEGORIES.items():
        if any(re.search(rf"\b{re.escape(w)}", n) for w in words):
            return cat
    return "other"


def _add(p: dict) -> str:
    amount = st.to_float(p.get("amount"))
    if not amount or amount <= 0:
        return "Sir, how much was it?"
    note = (p.get("note") or "").strip()
    cat = (p.get("category") or "").strip().lower()
    if cat not in CATEGORIES and cat != "other":
        cat = _guess_category(f"{note} {cat}")
    day = college.parse_date(p.get("date", "")) or date.today()
    with st.lock():
        data = st.load(NAME, DEFAULT)
        data["items"].append({"id": college.new_id(), "date": day.isoformat(), "amount": round(amount, 2),
                              "category": cat, "note": note})
        st.save(NAME, data)
    today_total = _total(_items("today"))
    msg = f"Logged {_rs(amount)} for {note or cat}. Today: {_rs(today_total)}."
    warn = _budget_check()
    return msg + (" " + warn if warn else "")


def _set_budget(p: dict) -> str:
    amount = st.to_float(p.get("amount"))
    if not amount or amount <= 0:
        return "Sir, what should the monthly budget be?"
    with st.lock():
        data = st.load(NAME, DEFAULT)
        data["budget"] = amount
        data["warned"] = {}
        st.save(NAME, data)
    spent = _total(_items("month"))
    return f"Monthly budget set to {_rs(amount)}. You've used {_rs(spent)} so far this month."


def _undo() -> str:
    with st.lock():
        data = st.load(NAME, DEFAULT)
        if not data["items"]:
            return "Nothing to undo, sir."
        it = data["items"].pop()
        st.save(NAME, data)
    return f"Removed {_rs(it['amount'])} for {it['note'] or it['category']}, sir."


# ── queries ──────────────────────────────────────────────────────────────────
def _range(period: str) -> tuple[date, date]:
    t = date.today()
    p = (period or "month").lower().replace(" ", "_")
    if p == "today":
        return t, t
    if p == "yesterday":
        return t - timedelta(days=1), t - timedelta(days=1)
    if p in ("week", "this_week", "last_7_days"):
        return t - timedelta(days=6), t
    if p == "last_month":
        end = t.replace(day=1) - timedelta(days=1)
        return end.replace(day=1), end
    return t.replace(day=1), t


def _items(period: str, category: str = "") -> list[dict]:
    a, b = _range(period)
    cat = category.lower().strip()
    return [i for i in st.load(NAME, DEFAULT)["items"]
            if a.isoformat() <= i["date"] <= b.isoformat() and (not cat or i["category"] == cat)]


def _total(items: list[dict]) -> float:
    return sum(i["amount"] for i in items)


def _by_cat(items: list[dict]) -> list[tuple[str, float]]:
    d = defaultdict(float)
    for i in items:
        d[i["category"]] += i["amount"]
    return sorted(d.items(), key=lambda x: -x[1])


def _summary(period: str, category: str) -> str:
    items = _items(period, category)
    label = {"today": "today", "yesterday": "yesterday", "week": "in the last 7 days",
             "last_month": "last month"}.get(period.lower().replace(" ", "_"), "this month")
    if not items:
        return f"No expenses {label}{' on ' + category if category else ''}, sir."
    msg = f"You spent {_rs(_total(items))} {label}{' on ' + category if category else ''}"
    if not category:
        msg += ": " + ", ".join(f"{c} {_rs(v)}" for c, v in _by_cat(items)[:4])
    msg += "."
    budget = st.load(NAME, DEFAULT)["budget"]
    if budget and label == "this month":
        left = budget - _total(_items("month"))
        msg += f" {_rs(left)} left of your {_rs(budget)} budget." if left >= 0 else \
            f" You're {_rs(-left)} over budget."
    return msg


def _budget_check() -> str:
    with st.lock():
        data = st.load(NAME, DEFAULT)
        budget = data["budget"]
        if not budget:
            return ""
        month = date.today().strftime("%Y-%m")
        spent = _total(_items("month"))
        warned = data["warned"].setdefault(month, [])
        for level in (100, 80):
            if spent >= budget * level / 100 and level not in warned:
                warned += [lv for lv in (80, 100) if lv <= level and lv not in warned]
                st.save(NAME, data)
                if level == 100:
                    msg = f"Heads-up: you've crossed your monthly budget — {_rs(spent)} of {_rs(budget)}."
                else:
                    msg = f"You've used {spent / budget:.0%} of this month's budget; {_rs(budget - spent)} left."
                bus.toast("Budget", msg)
                bus.phone("💸 " + msg)
                return msg
    return ""


def _export(p: dict) -> str:
    period = p.get("period") or "month"
    items = sorted(_items(period), key=lambda i: i["date"])
    if not items:
        return "No expenses to export for that period, sir."
    EXPORT_DIR.mkdir(exist_ok=True)
    a, _ = _range(period)
    base = EXPORT_DIR / f"expenses_{a.strftime('%Y-%m')}_{period}"
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses"
        ws.append(["Date", "Amount (₹)", "Category", "Note"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for i in items:
            ws.append([i["date"], i["amount"], i["category"], i["note"]])
        ws.append([])
        ws.append(["Total", f"=SUM(B2:B{len(items) + 1})"])
        ws[f"A{len(items) + 3}"].font = Font(bold=True)
        s = wb.create_sheet("By category")
        s.append(["Category", "Amount (₹)"])
        for cat, v in _by_cat(items):
            s.append([cat, v])
        for sheet, widths in ((ws, (12, 12, 12, 40)), (s, (14, 12))):
            for col, w in zip("ABCD", widths):
                sheet.column_dimensions[col].width = w
        path = base.with_suffix(".xlsx")
        wb.save(path)
    except ImportError:
        import csv
        path = base.with_suffix(".csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["date", "amount", "category", "note"])
            for i in items:
                w.writerow([i["date"], i["amount"], i["category"], i["note"]])
    try:
        if platform.system() == "Windows":
            os.startfile(path)            # noqa: opens in Excel
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass
    return f"Exported {st.plural(len(items), 'expense')} to exports/{path.name}, sir."


def call_facts(kind: str) -> list[str]:
    data = st.load(NAME, DEFAULT)
    if not data["items"] and not data["budget"]:
        return []
    month = _total(_items("month"))
    budget_line = ""
    if data["budget"]:
        budget_line = f"month so far {_rs(month)} of {_rs(data['budget'])} budget"
    if kind == "weekly":
        w = _items("week")
        top = _by_cat(w)[:3]
        return [f"Spending last 7 days: {_rs(_total(w))}" + (f" (top: {', '.join(f'{c} {_rs(v)}' for c, v in top)})" if top else "")
                + (f"; {budget_line}" if budget_line else "")]
    if kind == "recap":
        t = _items("today")
        line = f"Spent today: {_rs(_total(t))}" if t else "No expenses logged today (ask if they spent anything)"
        return [line + (f"; {budget_line}" if budget_line else "")]
    if data["budget"]:
        days_left = (date.today().replace(day=28) + timedelta(days=4)).replace(day=1) - date.today()
        left = data["budget"] - month
        if left > 0:
            return [f"Budget: {_rs(left)} left this month, about {_rs(round(left / max(1, days_left.days)))} a day"]
        return [f"Budget: already {_rs(-left)} over this month"]
    return []
